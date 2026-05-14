import os
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from migration.export._writer import ChunkedXlsx
from migration.export import exporters
from migration.export import runner as export_runner
from migration.logging_setup import setup_logging


@pytest.fixture
def logger(tmp_path):
    return setup_logging(str(tmp_path / 'logs'), verbose=False)


@pytest.fixture
def mock_odoo():
    o = MagicMock()
    o.safe_search_read.return_value = []
    o.read_group.return_value = []
    return o


def test_chunked_xlsx_splits_at_chunk_size(tmp_path):
    w = ChunkedXlsx(str(tmp_path), 'things', ['A', 'B'], chunk_size=2,
                    prefix='00_')
    for i in range(5):
        w.write_row({'A': i, 'B': f'b{i}'})
    w.close()
    files = sorted(os.path.basename(p) for p in w.paths)
    assert files == ['00_things_001.xlsx', '00_things_002.xlsx',
                     '00_things_003.xlsx']
    # 2 + 2 + 1 = 5 rows total
    wb = load_workbook(w.paths[0])
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == ('A', 'B')
    assert len(rows) == 3  # header + 2
    wb = load_workbook(w.paths[2])
    rows = list(wb.active.iter_rows(values_only=True))
    assert len(rows) == 2  # header + 1


def test_chunked_xlsx_no_rows_writes_nothing(tmp_path):
    w = ChunkedXlsx(str(tmp_path), 'empty', ['A'])
    w.close()
    assert w.paths == []
    assert w.total_rows == 0


def test_export_chart_of_accounts_emits_correct_columns(tmp_path, mock_odoo,
                                                        logger):
    mock_odoo.safe_search_read.side_effect = [
        [{'id': 1, 'name': 'Sales', 'code': '4000',
          'account_type': 'income', 'currency_id': [1, 'USD']},
         {'id': 2, 'name': 'Cash', 'code': '1000',
          'account_type': 'asset_cash', 'currency_id': [1, 'USD']}],
        [],
    ]
    summary = exporters.export_chart_of_accounts(
        mock_odoo, str(tmp_path), 5000, logger,
    )
    assert summary['rows'] == 2
    wb = load_workbook(summary['files'][0])
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == tuple(exporters.ACCOUNT_COLUMNS)
    # Find the Sales row
    sales = next(r for r in rows[1:] if r[1] == '4000')
    assert sales[0].startswith('4000 Sales [Odoo #1]')
    assert sales[2] == 'Income'  # account_type label
    assert sales[5] == 'odoo:1'  # reference number


def test_export_contacts_splits_customer_and_vendor(tmp_path, mock_odoo,
                                                    logger):
    # No self-partner.
    def side_effect(model, domain, fields=None, **kwargs):
        if model == 'res.company':
            return [{'id': 1, 'partner_id': False}]
        if model == 'res.partner':
            if domain and isinstance(domain[0], str) and domain[0] == '|':
                # Variant pre-scan.
                return [{'id': 7, 'customer_rank': 1, 'supplier_rank': 0},
                        {'id': 8, 'customer_rank': 0, 'supplier_rank': 1},
                        {'id': 9, 'customer_rank': 1, 'supplier_rank': 1}]
            return [
                {'id': 7, 'name': 'Acme', 'company_name': 'Acme Inc',
                 'is_company': True, 'parent_id': False,
                 'email': 'a@a.com', 'phone': '1', 'mobile': '',
                 'website': '', 'vat': '',
                 'street': '1 Main', 'street2': '', 'city': 'C',
                 'zip': 'Z', 'state_id': False, 'country_id': False,
                 'currency_id': False, 'customer_rank': 1, 'supplier_rank': 0,
                 'active': True},
                {'id': 8, 'name': 'Vendy', 'company_name': 'Vendy LLC',
                 'is_company': True, 'parent_id': False,
                 'email': '', 'phone': '', 'mobile': '', 'website': '',
                 'vat': '', 'street': '', 'street2': '', 'city': '',
                 'zip': '', 'state_id': False, 'country_id': False,
                 'currency_id': False, 'customer_rank': 0, 'supplier_rank': 1,
                 'active': True},
                {'id': 9, 'name': 'Both', 'company_name': 'Both Co',
                 'is_company': True, 'parent_id': False,
                 'email': '', 'phone': '', 'mobile': '', 'website': '',
                 'vat': '', 'street': '', 'street2': '', 'city': '',
                 'zip': '', 'state_id': False, 'country_id': False,
                 'currency_id': False, 'customer_rank': 1, 'supplier_rank': 1,
                 'active': True},
            ]
        return []

    mock_odoo.safe_search_read.side_effect = side_effect
    summary = exporters.export_contacts(mock_odoo, str(tmp_path), 5000, logger)
    assert summary['customers'] == 2  # Acme + Both
    assert summary['vendors'] == 2  # Vendy + Both

    # Make sure files landed with the right prefixes.
    fnames = [os.path.basename(p) for p in summary['files']]
    assert any(f.startswith('03_customers_') for f in fnames)
    assert any(f.startswith('04_vendors_') for f in fnames)


def test_export_contacts_skips_self_partner(tmp_path, mock_odoo, logger):
    def side_effect(model, domain, fields=None, **kwargs):
        if model == 'res.company':
            return [{'id': 1, 'partner_id': [42, 'Our Co']}]
        if model == 'res.partner':
            if domain and isinstance(domain[0], str) and domain[0] == '|':
                return [{'id': 42, 'customer_rank': 5, 'supplier_rank': 0},
                        {'id': 7, 'customer_rank': 1, 'supplier_rank': 0}]
            return [{'id': 7, 'name': 'Acme', 'company_name': '',
                     'is_company': True, 'parent_id': False,
                     'email': '', 'phone': '', 'mobile': '', 'website': '',
                     'vat': '', 'street': '', 'street2': '', 'city': '',
                     'zip': '', 'state_id': False, 'country_id': False,
                     'currency_id': False, 'customer_rank': 1,
                     'supplier_rank': 0, 'active': True}]
        return []

    mock_odoo.safe_search_read.side_effect = side_effect
    summary = exporters.export_contacts(mock_odoo, str(tmp_path), 5000, logger)
    assert summary['self_partner_skipped'] == 42
    assert summary['customers'] == 1  # only Acme, not the self-partner


def test_export_runner_writes_import_guide(tmp_path, mock_odoo, logger):
    # Empty Odoo — every exporter is a no-op but the guide should still write.
    export_runner.run(mock_odoo, str(tmp_path), 5000, logger,
                      phases=['chart_of_accounts'])
    guide = tmp_path / 'IMPORT_GUIDE.md'
    assert guide.exists()
    text = guide.read_text()
    assert 'Chart of Accounts' in text
    assert 'Customers' in text
    assert 'Foreign-key matching tips' in text


def test_export_runner_rejects_unknown_phase(tmp_path, mock_odoo, logger):
    with pytest.raises(ValueError, match='Unknown export phase'):
        export_runner.run(mock_odoo, str(tmp_path), 5000, logger,
                          phases=['bogus'])


def test_export_time_entries_keeps_employee_in_description(tmp_path, mock_odoo,
                                                           logger):
    def side_effect(model, domain, fields=None, **kwargs):
        if model == 'project.project':
            return [{'id': 50, 'name': 'Big Build',
                     'analytic_account_id': [777, 'BB Analytic']}]
        if model == 'account.analytic.account':
            return []
        if model == 'account.analytic.line':
            return [{'id': 1, 'name': 'Welded beam',
                     'date': '2024-05-01', 'unit_amount': 4.5,
                     'project_id': [50, 'Big Build'],
                     'account_id': [777, 'BB Analytic'],
                     'employee_id': [3, 'John Smith'],
                     'user_id': [9, 'admin']}]
        return []

    mock_odoo.safe_search_read.side_effect = side_effect
    summary = exporters.export_time_entries(
        mock_odoo, str(tmp_path), 5000, logger,
    )
    assert summary['rows'] == 1
    wb = load_workbook(summary['files'][0])
    rows = list(wb.active.iter_rows(values_only=True))
    header = rows[0]
    data = dict(zip(header, rows[1]))
    assert data['Hours'] == 4.5
    assert data['Project Name'].startswith('Big Build [Odoo #50]')
    assert '[John Smith]' in data['Description']
